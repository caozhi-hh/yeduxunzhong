# tools/export.py - Word 导出工具（精美排版版）
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
import re
import emoji

# 主题色
GREEN_DARK = (0x1B, 0x5E, 0x20)    # 深绿
GREEN_MAIN = (0x2E, 0x7D, 0x32)     # 主绿
GREEN_LIGHT = (0x4C, 0xAF, 0x50)    # 浅绿
GREEN_BG = (0xE8, 0xF5, 0xE9)       # 背景绿
ORANGE = (0xE6, 0x51, 0x00)         # 强调橙
GRAY_TEXT = (0x61, 0x61, 0x61)      # 正文灰
GRAY_LIGHT = (0xBD, 0xBD, 0xBD)     # 浅灰
WHITE = (0xFF, 0xFF, 0xFF)


def _strip_emoji(text: str) -> str:
    return emoji.replace_emoji(text, replace="")


def _clean_md(text: str) -> str:
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def _set_cell_shading(cell, color_hex: str):
    shading = cell._element.get_or_add_tcPr()
    elem = shading.makeelement(qn('w:shd'), {qn('w:fill'): color_hex, qn('w:val'): 'clear'})
    shading.append(elem)


def _set_cell_border(cell, **kwargs):
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge, val in kwargs.items():
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), val.get('val', 'single'))
        element.set(qn('w:sz'), val.get('sz', '4'))
        element.set(qn('w:color'), val.get('color', '000000'))
        element.set(qn('w:space'), val.get('space', '0'))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def _add_run(paragraph, text, size=11, bold=False, color=None, font_name="微软雅黑"):
    run = paragraph.add_run(_strip_emoji(str(text)))
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    return run


def _set_page_margins(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def export_to_word(itinerary: dict, filename: str = None) -> str:
    title_text = _strip_emoji(itinerary.get("title", "旅行攻略"))
    if not filename:
        filename = f"{title_text}.docx"

    doc = Document()
    _set_page_margins(doc)

    # 全局样式
    style = doc.styles['Normal']
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(*GRAY_TEXT)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.4

    # ========== 封面区 ==========
    # 顶部装饰线
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_run(p, "", size=2)
    # 用绿色横线模拟
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '24')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '2E7D32')
    pBdr.append(bottom)
    p._element.get_or_add_pPr().append(pBdr)

    # 大标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(8)
    _add_run(p, title_text, size=32, bold=True, color=GREEN_DARK)

    # 副标题
    dep = itinerary.get("dep_date", "")
    ret = itinerary.get("ret_date", "")
    budget_total = itinerary.get("budget_total", "")
    user_type = itinerary.get("user_type", "")
    info_parts = []
    if dep or ret:
        info_parts.append(f"{dep} ~ {ret}")
    if user_type:
        info_parts.append(user_type)
    if budget_total:
        info_parts.append(f"预算 {budget_total} 元/人")

    if info_parts:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(6)
        _add_run(p2, "  |  ".join(info_parts), size=11, color=GRAY_LIGHT)

    # 底部装饰线
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pBdr2 = OxmlElement('w:pBdr')
    top2 = OxmlElement('w:top')
    top2.set(qn('w:val'), 'single')
    top2.set(qn('w:sz'), '12')
    top2.set(qn('w:space'), '1')
    top2.set(qn('w:color'), '4CAF50')
    pBdr2.append(top2)
    p._element.get_or_add_pPr().append(pBdr2)

    doc.add_paragraph("")

    # ========== 正文 ==========
    overview = itinerary.get("overview", "")
    if overview:
        lines = overview.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            if not line:
                i += 1
                continue

            # ## 二级标题 → 绿色底色标题
            if line.startswith("## "):
                heading_text = _strip_emoji(line.replace("## ", "").strip())
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                # 绿色背景效果
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E8F5E9')
                shading.set(qn('w:val'), 'clear')
                p._element.get_or_add_pPr().append(shading)
                _add_run(p, f"  {heading_text}", size=16, bold=True, color=GREEN_MAIN)
                i += 1
                continue

            # ### 三级标题
            if line.startswith("### "):
                heading_text = _strip_emoji(line.replace("### ", "").strip())
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                # 左侧绿色竖线
                pBdr3 = OxmlElement('w:pBdr')
                left3 = OxmlElement('w:left')
                left3.set(qn('w:val'), 'single')
                left3.set(qn('w:sz'), '24')
                left3.set(qn('w:space'), '4')
                left3.set(qn('w:color'), '2E7D32')
                pBdr3.append(left3)
                p._element.get_or_add_pPr().append(pBdr3)
                _add_run(p, f"  {heading_text}", size=13, bold=True, color=(0x33, 0x33, 0x33))
                i += 1
                continue

            # #### 四级标题
            if line.startswith("#### "):
                heading_text = _strip_emoji(line.replace("#### ", "").strip())
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                _add_run(p, heading_text, size=11, bold=True, color=GREEN_MAIN)
                i += 1
                continue

            # 分隔线
            if line == "---":
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pBdr_line = OxmlElement('w:pBdr')
                bottom_line = OxmlElement('w:bottom')
                bottom_line.set(qn('w:val'), 'single')
                bottom_line.set(qn('w:sz'), '6')
                bottom_line.set(qn('w:space'), '1')
                bottom_line.set(qn('w:color'), 'E0E0E0')
                pBdr_line.append(bottom_line)
                p._element.get_or_add_pPr().append(pBdr_line)
                i += 1
                continue

            # 表格
            if line.startswith("|") and "|" in line[1:]:
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    stripped = lines[i].strip()
                    if not stripped.startswith("|-"):
                        cells = [c.strip() for c in stripped.split("|")[1:-1]]
                        table_lines.append(cells)
                    i += 1

                if table_lines:
                    num_cols = len(table_lines[0])
                    table = doc.add_table(rows=0, cols=num_cols)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    for row_idx, row_data in enumerate(table_lines):
                        row = table.add_row()
                        for col_idx, cell_text in enumerate(row_data):
                            if col_idx < num_cols:
                                cell = row.cells[col_idx]
                                is_header = (row_idx == 0)
                                cell.text = ""
                                p = cell.paragraphs[0]
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                _add_run(p, _strip_emoji(_clean_md(cell_text)),
                                        size=9 if not is_header else 9,
                                        bold=is_header,
                                        color=WHITE if is_header else GRAY_TEXT)
                                if is_header:
                                    _set_cell_shading(cell, "2E7D32")
                                elif row_idx % 2 == 0:
                                    _set_cell_shading(cell, "F1F8E9")

                    # 设置列宽自适应
                    for row in table.rows:
                        for cell in row.cells:
                            cell.width = Cm(max(2.5, 16 / num_cols))

                    doc.add_paragraph("")
                continue

            # 加粗开头段落
            bold_match = re.match(r'^\*\*(.+?)\*\*(.*)', line)
            if bold_match:
                p = doc.add_paragraph()
                _add_run(p, _strip_emoji(bold_match.group(1)), size=10.5, bold=True, color=(0x33, 0x33, 0x33))
                rest = bold_match.group(2)
                if rest:
                    _add_run(p, _strip_emoji(_clean_md(rest)), size=10.5, color=GRAY_TEXT)
                i += 1
                continue

            # 列表项（- 开头）
            if line.startswith("- ") or line.startswith("* "):
                content = _strip_emoji(_clean_md(line[2:]))
                p = doc.add_paragraph()
                _add_run(p, "  ", size=10.5)
                _add_run(p, content, size=10.5, color=GRAY_TEXT)
                i += 1
                continue

            # 普通段落
            content = _strip_emoji(_clean_md(line))
            if content.strip():
                p = doc.add_paragraph()
                _add_run(p, content, size=10.5, color=GRAY_TEXT)
            i += 1

    # ========== 页脚 ==========
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pBdr_foot = OxmlElement('w:pBdr')
    top_foot = OxmlElement('w:top')
    top_foot.set(qn('w:val'), 'single')
    top_foot.set(qn('w:sz'), '6')
    top_foot.set(qn('w:space'), '4')
    top_foot.set(qn('w:color'), 'E0E0E0')
    pBdr_foot.append(top_foot)
    p._element.get_or_add_pPr().append(pBdr_foot)
    _add_run(p, "野渡寻踪 | AI 旅行攻略", size=8, color=GRAY_LIGHT)

    # 保存
    output_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath
