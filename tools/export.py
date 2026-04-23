# tools/export.py - Word 导出工具（简约大气版）
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os
import re
from io import BytesIO


def _set_cell_shading(cell, color_hex):
    """设置单元格背景色"""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear',
    })
    shading.append(shading_elem)


def _add_styled_paragraph(doc, text, font_size=11, bold=False, color=None, align=None, space_after=6):
    """添加带样式的段落"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def export_to_word(itinerary: dict, filename: str = None) -> str:
    """将攻略导出为简约大气的 Word 文档。"""
    title_text = itinerary.get("title", "旅行攻略")
    if not filename:
        filename = f"{title_text}.docx"

    doc = Document()

    # 全局样式
    style = doc.styles['Normal']
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.3

    # ========== 封面区 ==========
    # 大标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    run = p.add_run("🌿 " + title_text)
    run.font.size = Pt(28)
    run.bold = True
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")

    # 副标题信息
    dep = itinerary.get("dep_date", "")
    ret = itinerary.get("ret_date", "")
    budget_total = itinerary.get("budget_total", "")
    user_type = itinerary.get("user_type", "")
    if dep or ret:
        sub = f"{dep} ~ {ret}"
        if user_type:
            sub += f"  |  {user_type}"
        if budget_total:
            sub += f"  |  预算 {budget_total} 元"
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(sub)
        run2.font.size = Pt(12)
        run2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph("")  # 空行

    # ========== 正文：直接渲染 Markdown 攻略内容 ==========
    overview = itinerary.get("overview", "")
    if overview:
        # 按 Markdown 结构解析并渲染
        lines = overview.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            # 跳过空行
            if not line:
                i += 1
                continue

            # 二级标题 ## Day X
            if line.startswith("## "):
                heading_text = line.replace("## ", "").strip()
                h = doc.add_heading(heading_text, level=2)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(0xE8, 0x6B, 0x30)
                i += 1
                continue

            # 三级标题 ###
            if line.startswith("### "):
                heading_text = line.replace("### ", "").strip()
                h = doc.add_heading(heading_text, level=3)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                i += 1
                continue

            # 四级标题 ####
            if line.startswith("#### "):
                heading_text = line.replace("#### ", "").strip()
                _add_styled_paragraph(doc, heading_text, font_size=12, bold=True,
                                     color=(0x55, 0x55, 0x55))
                i += 1
                continue

            # 分隔线
            if line == "---":
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run("· · · · · · · · · ·")
                run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
                run.font.size = Pt(10)
                i += 1
                continue

            # 表格检测
            if line.startswith("|") and "|" in line[1:]:
                # 收集表格所有行
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    stripped = lines[i].strip()
                    if not stripped.startswith("|-"):  # 跳过分隔行
                        cells = [c.strip() for c in stripped.split("|")[1:-1]]
                        table_lines.append(cells)
                    i += 1

                if table_lines:
                    num_cols = len(table_lines[0])
                    table = doc.add_table(rows=0, cols=num_cols)
                    table.style = "Light Grid Accent 1"
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER

                    for row_idx, row_data in enumerate(table_lines):
                        row = table.add_row()
                        for col_idx, cell_text in enumerate(row_data):
                            if col_idx < num_cols:
                                cell = row.cells[col_idx]
                                cell.text = _clean_md(cell_text)
                                for paragraph in cell.paragraphs:
                                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                    for run in paragraph.runs:
                                        run.font.size = Pt(10)
                                        run.font.name = "微软雅黑"
                                        run._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")
                    doc.add_paragraph("")
                continue

            # 加粗段落（如 **07:00 - 07:30 | 早餐**）
            bold_match = re.match(r'^\*\*(.+?)\*\*(.*)', line)
            if bold_match:
                p = doc.add_paragraph()
                run = p.add_run(bold_match.group(1))
                run.bold = True
                run.font.size = Pt(11)
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")
                rest = bold_match.group(2)
                if rest:
                    run2 = p.add_run(_clean_md(rest))
                    run2.font.size = Pt(11)
                    run2.font.name = "微软雅黑"
                    run2._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")
                i += 1
                continue

            # 普通段落
            _add_styled_paragraph(doc, _clean_md(line), font_size=11)
            i += 1

        # 在文档末尾插入景点图片
        spot_images = itinerary.get("spot_images", {})
        if spot_images:
            doc.add_page_break()
            h = doc.add_heading("📸 景点一览", level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0xE8, 0x6B, 0x30)

            cols_per_row = 2
            spot_list = [(name, path) for name, path in spot_images.items() if path]
            for i in range(0, len(spot_list), cols_per_row):
                row_spots = spot_list[i:i+cols_per_row]
                # 每张图片占一段，居中
                for name, path in row_spots:
                    try:
                        if os.path.exists(path):
                            p = doc.add_paragraph()
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            run = p.add_run(name)
                            run.bold = True
                            run.font.size = Pt(10)
                            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                            doc.add_picture(path, width=Cm(7))
                            last_p = doc.paragraphs[-1]
                            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except:
                        pass

    # ========== 页脚 ==========
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("— 野渡寻踪 · AI 旅行攻略 —")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "微软雅黑")

    # 保存
    output_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath


def _clean_md(text):
    """清理 Markdown 格式符号"""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text
